"""No-hardware checks for the P3--P6 workflow pages."""

import os
import gc
import sqlite3
import tempfile
import unittest
from datetime import datetime
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt6.QtWidgets import QApplication, QFrame, QLineEdit, QTabWidget, QVBoxLayout, QWidget
from PyQt6.QtGui import QColor, QImage, QPainter
from PyQt6 import sip

from public.component.system_status import SystemStatusPage
from public.component.workflow_pages import (
    DataAnalysisPage,
    DataViewPage,
    ExperimentControlPage,
    MachineVisionPage,
    TrajectoryComparisonPlot,
)


class WorkflowPageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    @staticmethod
    def context():
        return {
            "experiment_status": "not_started",
            "elapsed_seconds": None,
            "health_state": "unknown",
            "experiment_db_path": None,
        }

    @staticmethod
    def health():
        return {"components": [], "alerts": []}

    def test_p3_has_start_stop_and_calibration_without_pause(self):
        page = ExperimentControlPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.deleteLater)
        self.assertEqual(page.start_button.text(), "▶ 开始实验")
        self.assertEqual(page.stop_button.text(), "■ 终止实验")
        self.assertFalse(any("暂停" in button.text() for button in page.findChildren(type(page.start_button))))
        self.assertEqual(
            set(page.calibration_buttons),
            {"air", "zero", "span", "all", "stop_zero", "stop_span", "stop_all"},
        )
        self.assertEqual(page.calibration_buttons["air"].text(), "开始空气校准")
        self.assertEqual(page.calibration_details_button.text(), "显示校准信息详情")
        self.assertTrue(hasattr(page, "main_log_list"))
        self.assertGreaterEqual(page.main_log_list.columnWidth(0), 96)
        self.assertGreaterEqual(page.main_log_list.columnWidth(1), 68)

    def test_p3_summary_cards_are_equal_and_duration_refreshes_each_second(self):
        page = ExperimentControlPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)

        self.assertEqual(page._refresh_timer.interval(), 1000)
        self.assertEqual(len(page.summary_cards), 4)
        self.assertEqual(
            {card.minimumHeight() for card in page.summary_cards},
            {118},
        )
        self.assertEqual(
            [page.status_grid.columnStretch(column) for column in range(4)],
            [1, 1, 1, 1],
        )

    def test_p3_calibration_requires_all_three_air_modules(self):
        from public.config_class.global_setting import global_setting

        context = lambda: {**self.context(), "experiment_status": "running"}
        main_window = Mock()
        page = ExperimentControlPage(main_window=main_window, context_provider=context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)
        with patch.dict(global_setting._setting, {
            "environment_module_only": False, "air_modules_all_valid": False,
            "send_message_queue": Mock(),
        }):
            page.refresh()
            self.assertFalse(page.calibration_buttons["zero"].isEnabled())
            self.assertFalse(page.calibration_buttons["span"].isEnabled())
            page._start_calibration("zero")
            self.assertIn("UFC、UGC、ZOS", page.operation_hint.text())
            with patch.dict(global_setting._setting, {"air_modules_all_valid": True}):
                page.refresh()
                self.assertTrue(page.calibration_buttons["zero"].isEnabled())
                self.assertTrue(page.calibration_buttons["span"].isEnabled())

    def test_p3_air_calibration_does_not_start_without_detection(self):
        from public.config_class.global_setting import global_setting

        main_window = Mock()
        page = ExperimentControlPage(main_window=main_window, context_provider=self.context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)
        with patch.dict(global_setting._setting, {
            "environment_module_only": False, "air_modules_all_valid": False,
        }):
            page.refresh()
            self.assertFalse(page.calibration_buttons["air"].isEnabled())
            page._start_air_calibration()
            main_window.start_experiment.assert_not_called()
            self.assertIn("UFC、UGC、ZOS", page.operation_hint.text())

    def test_p3_progress_is_in_status_card_and_survives_refresh(self):
        context = lambda: {**self.context(), "experiment_status": "starting"}
        page = ExperimentControlPage(context_provider=context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)

        self.assertTrue(page.status_card.isAncestorOf(page.operation_progress))
        self.assertTrue(page.status_card.isAncestorOf(page.operation_label))
        page.set_operation_progress(37, "空气校准进行中")
        page.refresh()
        self.assertEqual(page.operation_progress.value(), 37)
        self.assertIn("37%", page.progress_title.text())
        self.assertIn("空气校准进行中", page.operation_label.text())

    def test_p3_shows_start_and_stop_countdown_in_operation_area(self):
        context = {
            **self.context(),
            "experiment_status": "starting",
            "operation_countdown_phase": "启动",
            "operation_countdown_seconds": 58,
        }
        page = ExperimentControlPage(context_provider=lambda: context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)

        page.set_operation_progress(12, "正在启动气路")
        page.refresh()
        self.assertIn("启动剩余时间：58秒", page.operation_label.text())

        context.update({
            "experiment_status": "stopping",
            "operation_countdown_phase": "停止",
            "operation_countdown_seconds": 21,
        })
        page.set_operation_progress(40, "正在停止实验")
        page.refresh()
        self.assertIn("停止剩余时间：21秒", page.operation_label.text())

    def test_p3_start_button_calls_main_window_and_gives_immediate_feedback(self):
        main_window = Mock()
        page = ExperimentControlPage(main_window=main_window, context_provider=self.context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)

        page.start_button.click()
        main_window.start_experiment.assert_called_once_with()
        self.assertIn("正在检查", page.operation_label.text())
        self.assertIn("后台连接", page.operation_hint.text())

    def test_p3_can_stop_partial_start_after_backend_failure(self):
        context = lambda: {
            **self.context(),
            "experiment_status": "start_failed",
            "can_stop_experiment": True,
        }
        page = ExperimentControlPage(context_provider=context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)

        self.assertFalse(page.start_button.isEnabled())
        self.assertTrue(page.stop_button.isEnabled())

    def test_p3_disables_repeat_stop_without_locking_the_page(self):
        context = lambda: {
            **self.context(),
            "experiment_status": "stopping",
            "can_stop_experiment": True,
        }
        page = ExperimentControlPage(context_provider=context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)

        page.refresh()
        self.assertTrue(page.isEnabled())
        self.assertFalse(page.start_button.isEnabled())
        self.assertFalse(page.stop_button.isEnabled())

    def test_p1_and_p3_share_duration_and_health_summary(self):
        context = lambda: {
            **self.context(),
            "experiment_status": "starting",
            "elapsed_seconds": 2402,
            "health_state": "normal",
            "enabled_channels": ["CH01"],
        }
        health = lambda: {
            "components": [
                {
                    "component": "monitor_collection",
                    "scope": "CH01",
                    "status": "normal",
                    "observed_at": 1_700_000_000,
                }
            ],
            "alerts": [],
        }
        p1 = SystemStatusPage(context_provider=context, health_source=health, refresh_interval_ms=60_000)
        p3 = ExperimentControlPage(context_provider=context, health_source=health)
        self.addCleanup(p1.stop_refresh)
        self.addCleanup(p1.deleteLater)
        self.addCleanup(p3.deleteLater)

        p1.refresh()
        p3.refresh()

        self.assertEqual(p1._overview_values["run_duration"].text(), "0天0小时40分2秒")
        self.assertEqual(p3.duration_value.text(), p1._overview_values["run_duration"].text())
        self.assertEqual(p3.health_value.text(), p1._overview_values["system_health"].text())
        self.assertNotEqual(p3.health_value.text(), "fault")

    def test_p4_without_database_is_explicitly_empty(self):
        page = DataViewPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.deleteLater)
        page.refresh()
        self.assertIn("暂无", page.status_label.text())
        self.assertEqual(page.table.rowCount(), 0)
        self.assertEqual(len(page.channel_filter_buttons), 9)

    def test_gui_database_directory_matches_monitor_storage_layout(self):
        from index.MainWindow_index import MainWindow_Index

        template_path = os.path.join(os.getcwd(), "1.template")
        self.assertTrue(os.path.isfile(template_path))
        config = {"STORAGE": {"fold_path": "./data/", "sub_fold_path": "/monitor_data/"}}
        def setting(key, default=None):
            return {
                "experiment_setting_file": template_path,
                "monitor_data": config,
                "start_experiment_time": 1799976000,
            }.get(key, default)

        with patch("index.MainWindow_index.global_setting.get_setting", side_effect=setting), patch(
            "index.MainWindow_index.time_util.get_format_file_from_time", return_value="2026_09_17_12_00_00_000"
        ):
            root = MainWindow_Index._get_experiment_root_path(None)
        expected = os.path.join(os.getcwd(), "data", "monitor_data", "1_2026_09_17_12_00_00_000")
        self.assertEqual(root, os.path.abspath(expected))

    def test_gui_database_directory_requires_start_time_and_matches_service_fallback(self):
        from index.MainWindow_index import MainWindow_Index

        config = {"STORAGE": {"fold_path": "./data/", "sub_fold_path": "/monitor_data/"}}
        settings = {
            "experiment_setting_file": None,
            "monitor_data": config,
            "start_experiment_time": None,
        }
        with patch("index.MainWindow_index.global_setting.get_setting", side_effect=lambda key, default=None: settings.get(key, default)), patch(
            "index.MainWindow_index.time_util.get_format_file_from_time", return_value="run_start"
        ):
            self.assertIsNone(MainWindow_Index._get_experiment_root_path(None))
            settings["start_experiment_time"] = 1799976000
            self.assertEqual(
                MainWindow_Index._get_experiment_root_path(None),
                os.path.abspath(os.path.join(os.getcwd(), "data", "monitor_data", "_run_start")),
            )

    def test_gui_uses_the_single_experiment_directory_fixed_at_start(self):
        from index.MainWindow_index import MainWindow_Index

        experiment_path = os.path.abspath(os.path.join("data", "monitor_data", "_fixed_at_start"))
        with patch(
            "index.MainWindow_index.global_setting.get_setting",
            side_effect=lambda key, default=None: experiment_path if key == "experiment_root_path" else default,
        ):
            self.assertEqual(MainWindow_Index._get_experiment_root_path(None), experiment_path)

    def test_gui_export_root_is_the_project_not_the_process_working_directory(self):
        from index.MainWindow_index import MainWindow_Index

        with tempfile.TemporaryDirectory() as temporary_root:
            config = {"STORAGE": {"fold_path": "./data/", "sub_fold_path": "/monitor_data/"}}
            settings = {
                "experiment_setting_file": None,
                "monitor_data": config,
                "start_experiment_time": 1799976000,
            }
            with patch("index.MainWindow_index.os.getcwd", return_value=temporary_root), patch(
                "index.MainWindow_index.global_setting.get_setting",
                side_effect=lambda key, default=None: settings.get(key, default),
            ), patch(
                "index.MainWindow_index.time_util.get_format_file_from_time",
                return_value="run_start",
            ):
                resolved = MainWindow_Index._get_experiment_root_path(None)

        self.assertEqual(
            resolved,
            os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data", "monitor_data", "_run_start")),
        )
        self.assertFalse(resolved.startswith(os.path.abspath(temporary_root)))

    def test_gui_export_root_uses_the_user_selected_storage_location(self):
        from index.MainWindow_index import MainWindow_Index

        with tempfile.TemporaryDirectory() as selected_root:
            config = {"STORAGE": {"fold_path": "./data/", "sub_fold_path": "/monitor_data/"}}
            experiment_setting = Mock(storage_path=selected_root)
            settings = {
                "experiment_setting_file": None,
                "experiment_setting": experiment_setting,
                "experiment_setting_new": experiment_setting,
                "monitor_data": config,
                "start_experiment_time": 1799976000,
            }
            with patch(
                "index.MainWindow_index.global_setting.get_setting",
                side_effect=lambda key, default=None: settings.get(key, default),
            ), patch(
                "index.MainWindow_index.time_util.get_format_file_from_time",
                return_value="run_start",
            ):
                resolved = MainWindow_Index._get_experiment_root_path(None)

        self.assertEqual(resolved, os.path.abspath(os.path.join(selected_root, "_run_start")))

    def test_p4_navigation_reuses_page_and_keeps_channel_selection(self):
        from index.MainWindow_index import MainWindow_Index

        class WindowHarness:
            _cached_workflow_page = MainWindow_Index._cached_workflow_page
            _show_workflow_page = MainWindow_Index._show_workflow_page
            show_data_view_page = MainWindow_Index.show_data_view_page

            def __init__(self):
                self.tab_widget = QTabWidget()
                self.workflow_shell = Mock()
                self._system_status_context = lambda: {"experiment_db_path": None}
                self._system_health_view = lambda: {}

        window = WindowHarness()
        self.addCleanup(window.tab_widget.deleteLater)
        window.show_data_view_page()
        original = window.data_view_page
        original._select_channel_filter(3)
        another_page = QWidget()
        window._show_workflow_page(another_page, 1)
        self.assertFalse(original._refresh_timer.isActive())
        window.show_data_view_page()
        self.assertIs(window.data_view_page, original)
        self.assertEqual(window.tab_widget.count(), 2)
        self.assertEqual(original._selected_channel_filter, 3)
        self.assertTrue(original.channel_filter_buttons[3].isChecked())
        self.assertTrue(original._refresh_timer.isActive())
        original.stop_refresh()

    def test_all_six_top_navigation_steps_keep_page_inputs(self):
        from index.MainWindow_index import MainWindow_Index

        class Page(QWidget):
            def __init__(self, **_kwargs):
                super().__init__()
                self.field = QLineEdit(self)
                layout = QVBoxLayout(self)
                layout.addWidget(self.field)

            def start_refresh(self):
                pass

            def stop_refresh(self):
                pass

            def refresh(self):
                pass

            def on_calibration_event(self, _event):
                pass

        class WindowHarness:
            _cached_workflow_page = MainWindow_Index._cached_workflow_page
            _show_workflow_page = MainWindow_Index._show_workflow_page
            show_system_status_page = MainWindow_Index.show_system_status_page
            show_experiment_config_page = MainWindow_Index.show_experiment_config_page
            show_experiment_control_page = MainWindow_Index.show_experiment_control_page
            show_data_view_page = MainWindow_Index.show_data_view_page
            show_machine_vision_page = MainWindow_Index.show_machine_vision_page
            show_data_analysis_page = MainWindow_Index.show_data_analysis_page
            navigate_workflow_step = MainWindow_Index.navigate_workflow_step

            def __init__(self):
                self.tab_widget = QTabWidget()
                self.workflow_shell = Mock()
                self._system_status_context = lambda: {}
                self._system_health_view = lambda: {}
                self._active_calibration_channels = set()
                self._hide_legacy_module_windows = lambda: None

        window = WindowHarness()
        self.addCleanup(window.tab_widget.deleteLater)
        with patch.multiple(
            "index.MainWindow_index",
            SystemStatusPage=Page,
            ExperimentConfigPage=Page,
            ExperimentControlPage=Page,
            DataViewPage=Page,
            MachineVisionPage=Page,
            DataAnalysisPage=Page,
        ):
            pages = {}
            names = (
                "system_status_page", "experiment_config_page", "experiment_control_page",
                "data_view_page", "machine_vision_page", "data_analysis_page",
            )
            for step, name in enumerate(names, 1):
                window.navigate_workflow_step(step)
                pages[step] = getattr(window, name)
                pages[step].field.setText(f"saved-{step}")
            for step, name in reversed(list(enumerate(names, 1))):
                window.navigate_workflow_step(step)
                self.assertIs(getattr(window, name), pages[step])
                self.assertEqual(pages[step].field.text(), f"saved-{step}")
            self.assertEqual(window.tab_widget.count(), 6)

    def test_p4_open_during_startup_shows_real_rows_when_database_arrives(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            context = lambda: {
                "experiment_status": "starting",
                "experiment_db_path": db_path,
            }
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            self.assertIn("等待采集进程", page.status_label.text())
            self.assertEqual(page.table.rowCount(), 0)
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_cage_1 (id INTEGER, value REAL)")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (1, 42.0)")
            page.refresh()
            self.assertEqual(page.table.rowCount(), 1)
            self.assertIn("轮次汇总记录", page.status_label.text())

    def test_p4_refresh_finds_collection_database_by_current_start_time(self):
        from public.config_class.global_setting import global_setting

        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            monitor_root = os.path.join(temp_dir, "monitor_data")
            started = datetime(2026, 9, 19, 10, 20, 30, 456000)
            experiment_dir = os.path.join(
                monitor_root,
                "1_2026_09_19_10_20_30_456",
                "data",
            )
            os.makedirs(experiment_dir)
            db_path = os.path.join(experiment_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_cage_1 (id INTEGER, value REAL)")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (1, 42.0)")

            context_data = {
                "experiment_status": "starting",
                "started_at": None,
                "experiment_db_path": os.path.join(temp_dir, "predicted", "data.db"),
            }
            storage = {
                "STORAGE": {
                    "fold_path": temp_dir,
                    "sub_fold_path": "monitor_data",
                }
            }
            with patch.dict(global_setting._setting, {"monitor_data": storage, "start_experiment_time": None}):
                page = DataViewPage(
                    context_provider=lambda: dict(context_data),
                    health_source=self.health,
                )
                self.addCleanup(page.stop_refresh)
                self.addCleanup(page.deleteLater)
                self.assertEqual(page.table.rowCount(), 0)

                context_data["started_at"] = started.timestamp()
                page.reload_button.click()

                self.assertEqual(page.table.rowCount(), 1)
                self.assertEqual(page.table.item(0, 2).text(), "42.0")
                self.assertIn(os.path.abspath(db_path), page.status_label.text())

    def test_p4_export_current_query_uses_selected_save_path(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_cage_01 (id INTEGER, value REAL)")
                connection.execute("INSERT INTO Epoch_data_cage_01 VALUES (1, 11.0)")
                connection.commit()

            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.deleteLater)
            page.refresh()

            output_path = os.path.join(temp_dir, "selected-query.csv")
            with patch(
                "public.component.workflow_pages.QFileDialog.getSaveFileName",
                return_value=(output_path, "CSV (*.csv)"),
            ) as chooser:
                self.assertTrue(page.export_current_query())

            chooser.assert_called_once()
            self.assertEqual(page._last_export_directory, os.path.abspath(temp_dir))
            with open(output_path, "r", encoding="utf-8-sig", newline="") as stream:
                exported = stream.read()
            self.assertIn("01", exported)
            self.assertIn("11.0", exported)

    def test_p4_reads_epoch_rows_automatically_instead_of_raw_module_rows(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute('CREATE TABLE UFC_monitor_data_cage_1 (id INTEGER, time TEXT, flow_num REAL)')
                connection.execute('CREATE TABLE UFC_monitor_data_cage_2 (id INTEGER, time TEXT, flow_num REAL)')
                connection.execute('CREATE TABLE Epoch_data_cage_1 (id INTEGER, value REAL)')
                connection.execute("INSERT INTO UFC_monitor_data_cage_1 VALUES (1, '2026-09-17 12:00:00', 4.5)")
                connection.execute("INSERT INTO UFC_monitor_data_cage_2 VALUES (1, '2026-09-17 12:00:01', NULL)")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (1, 99.0)")

            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.deleteLater)
            page.refresh()
            self.assertEqual(page.table.rowCount(), 1)
            self.assertIn("轮次汇总记录", page.status_label.text())
            headers = [page.table.horizontalHeaderItem(index).text() for index in range(page.table.columnCount())]
            self.assertIn("value", headers)
            self.assertNotIn("flow_num", headers)
            page.channel_filter_buttons[1].click()
            self.assertEqual(page.table.rowCount(), 1)
            self.assertEqual(page.table.item(0, headers.index("value")).text(), "99.0")

    def test_p4_empty_tables_do_not_show_placeholder_columns(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_cage_1 (id INTEGER, value REAL)")
            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.deleteLater)
            page.refresh()
            self.assertEqual(page.table.rowCount(), 0)
            self.assertEqual(page.table.columnCount(), 0)

    def test_p4_shows_four_calculated_metrics_from_real_epoch_inputs(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute(
                    "CREATE TABLE Epoch_data_cage_1 (id INTEGER, UFC_flow_num REAL, "
                    "reference_dry_basis_oxygen_num REAL, reference_CO2_num REAL, "
                    "ZOS_dry_basis_oxygen_num REAL, UGC_CO2_num REAL)"
                )
                connection.execute(
                    "INSERT INTO Epoch_data_cage_1 VALUES (1, 1000, 20.93, 0.04, 20, 1)"
                )
            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            self.assertEqual(page.table.rowCount(), 1)  # no confirmation click
            headers = [page.table.horizontalHeaderItem(index).text() for index in range(page.table.columnCount())]
            for label in ("VO2氧耗量(L/min)", "VCO2二氧化碳产生量(L/min)", "RER呼吸熵", "EE能量消耗(kcal/hr)"):
                self.assertIn(label, headers)
                self.assertNotEqual(page.table.item(0, headers.index(label)).text(), "")

    def test_p4_all_channels_reads_aggregate_epoch_without_duplicate_cage_rows(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_all (id INTEGER, mouse_cage_number INTEGER, value REAL)")
                connection.execute("CREATE TABLE Epoch_data_cage_1 (id INTEGER, value REAL)")
                connection.execute("INSERT INTO Epoch_data_all VALUES (1, 1, 12.0)")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (1, 12.0)")
            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            self.assertEqual(page.table.rowCount(), 1)
            self.assertEqual(page.table.item(0, 0).text(), "1")

    def test_p4_displays_reference_cage_and_null_values_like_old_monitor(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_all (id INTEGER, mouse_cage_number INTEGER, value REAL)")
                connection.execute("INSERT INTO Epoch_data_all VALUES (1, 0, NULL)")
            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)

            headers = [page.table.horizontalHeaderItem(index).text() for index in range(page.table.columnCount())]
            self.assertEqual(page.table.rowCount(), 1)
            self.assertEqual(page.table.item(0, 0).text(), "参考笼")
            self.assertEqual(page.table.item(0, headers.index("鼠笼号")).text(), "参考笼")
            self.assertEqual(page.table.item(0, headers.index("value")).text(), "None")

    def test_p4_marks_entire_row_red_when_remarks_report_an_error(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute(
                    "CREATE TABLE Epoch_data_all "
                    "(id INTEGER, mouse_cage_number INTEGER, value REAL, remarks TEXT)"
                )
                connection.execute("INSERT INTO Epoch_data_all VALUES (1, 1, 10.0, NULL)")
                connection.execute("INSERT INTO Epoch_data_all VALUES (2, 2, NULL, '模块通信报错')")
            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataViewPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)

            self.assertEqual(page.table.rowCount(), 2)
            for column in range(page.table.columnCount()):
                self.assertEqual(page.table.item(0, column).foreground().color(), QColor(255, 0, 0))
                self.assertNotEqual(page.table.item(1, column).foreground().color(), QColor(255, 0, 0))

    def test_p4_canonical_epoch_uses_legacy_metadata_and_time_order(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_cage_1 (id INTEGER, time TEXT, mouse_cage_number INTEGER, WM_weight_num REAL)")
                connection.execute("CREATE TABLE Epoch_data_cage_1_meta (item_name TEXT, item_struct TEXT, description TEXT)")
                connection.execute("INSERT INTO Epoch_data_cage_1_meta VALUES ('WM_weight_num', 'REAL', '旧版称重值')")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (1, '2026-09-17 12:00:00', 1, 10.0)")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (2, '2026-09-17 12:01:00', 1, 11.0)")
            page = DataViewPage(
                context_provider=lambda: {**self.context(), "experiment_db_path": db_path},
                health_source=self.health,
            )
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            page.channel_filter_buttons[1].click()
            headers = [page.table.horizontalHeaderItem(index).text() for index in range(page.table.columnCount())]
            self.assertIn("旧版称重值", headers)
            self.assertEqual(page.table.rowCount(), 2)
            self.assertEqual(page.table.item(0, headers.index("旧版称重值")).text(), "11.0")

    def test_channel_buttons_bind_each_page_to_the_selected_channel(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE Epoch_data_cage_01 (id INTEGER, value REAL)")
                connection.execute("CREATE TABLE Epoch_data_cage_02 (id INTEGER, value REAL)")
                connection.execute("INSERT INTO Epoch_data_cage_01 VALUES (1, 11.0)")
                connection.execute("INSERT INTO Epoch_data_cage_02 VALUES (1, 22.0)")
                connection.execute("CREATE TABLE DetectionResults_cage_01 (x REAL, y REAL, z REAL, time REAL)")
                connection.execute("CREATE TABLE DetectionResults_cage_02 (x REAL, y REAL, z REAL, time REAL)")
                for index in range(4):
                    connection.execute("INSERT INTO DetectionResults_cage_01 VALUES (?, ?, ?, ?)", (0.0, 0.0, 0.0, index * 10.0))
                    connection.execute("INSERT INTO DetectionResults_cage_02 VALUES (?, ?, ?, ?)", (float(index), 0.0, 0.0, index * 10.0))
                connection.commit()

            context = lambda: {
                **self.context(),
                "experiment_db_path": db_path,
                "channel_individuals": {"CH01": {"id": "M001"}, "CH02": {"id": "M002"}},
            }

            data_page = DataViewPage(context_provider=context, health_source=self.health)
            data_page.channel_filter_buttons[2].click()  # CH02
            self.assertEqual(data_page._selected_channel_filter, 2)
            self.assertEqual(data_page.table.rowCount(), 1)
            self.assertIn("CH02", data_page.status_label.text())
            self.assertIn("22.0", data_page.table.item(0, 2).text())

            vision_page = MachineVisionPage(context_provider=context, health_source=self.health)
            vision_page.channel_buttons[1].click()  # CH02
            self.assertEqual(vision_page.selected_channel, 2)
            self.assertIn("CH02", vision_page.info_label.text())
            self.assertIn("M002", vision_page.info_label.text())
            self.assertEqual(len(vision_page._trajectory_records()), 4)

            analysis_page = DataAnalysisPage(context_provider=context, health_source=self.health)
            analysis_page.channel_buttons[1].click()  # CH02
            self.assertEqual(analysis_page.selected_channel, 2)
            self.assertIn("CH02", analysis_page.bio_label.text())
            self.assertIn("M002", analysis_page.bio_label.text())
            self.assertEqual(len(analysis_page._history_points), 4)
            data_page.stop_refresh()
            vision_page._frame_timer.stop()
            analysis_page.stop_refresh()
            data_page.deleteLater()
            vision_page.deleteLater()
            analysis_page.deleteLater()
            self.app.processEvents()
            gc.collect()

    def test_p5_without_camera_frame_does_not_claim_online(self):
        page = MachineVisionPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.deleteLater)
        page.refresh()
        self.assertIn("无新帧", page.frame_label.text())
        self.assertIn("未知", page.info_label.text())

    def test_p5_behavior_ui_has_no_confidence_field(self):
        page = MachineVisionPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.deleteLater)
        page.refresh()

        self.assertEqual(page.recognition_log.columnCount(), 3)
        self.assertNotIn("置信度", page.info_label.text())
        self.assertNotIn("置信度", page.recognition_count_label.text())
        headers = [page.recognition_log.horizontalHeaderItem(index).text() for index in range(3)]
        self.assertNotIn("置信度", headers)

    def test_p6_formulas_and_history_are_not_fabricated(self):
        page = DataAnalysisPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.deleteLater)
        page.refresh()
        self.assertEqual(page.metric_labels["night_day"].text(), "None / 公式待定")
        self.assertIn("CH01", page.bio_label.text())
        self.assertEqual(page.bio_metric_labels["weight"].text(), "未知")
        self.assertEqual(page.bio_metric_labels["metabolism"].text(), "None")

    def test_p6_compare_supports_multiple_channels_and_running_wheel(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                for channel in (1, 2):
                    connection.execute(
                        f"CREATE TABLE Epoch_data_cage_{channel} (time INTEGER, oxygen_num REAL, ENM_running_wheel_num REAL, vo2_consumption REAL, vco2_production REAL)"
                    )
                    connection.executemany(
                        f"INSERT INTO Epoch_data_cage_{channel} VALUES (?, ?, ?, ?, ?)",
                        [(1, 20.0 + channel, 2.0 + channel, 0.5 + channel / 10, 0.25 + channel / 20), (2, 21.0 + channel, 3.0 + channel, 0.6 + channel / 10, 0.3 + channel / 20)],
                    )
                connection.commit()

            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataAnalysisPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.deleteLater)
            page.left_metric.setCurrentIndex(5)  # running wheel
            page.left_channel_buttons[1].click()  # add CH02 to the left axis
            page.right_channel_buttons[1].click()  # add CH02 to the right axis

            self.assertEqual(page._left_channels, {1, 2})
            self.assertEqual(page._right_channels, {1, 2})
            self.assertEqual(len(page.compare_canvas.left_series), 2)
            self.assertEqual(page.compare_canvas.left_series[1][1], [4.0, 5.0])
            page.left_metric.setCurrentIndex(8)  # RER = VCO2 / VO2
            self.assertAlmostEqual(page.compare_canvas.left_series[0][1][0], 0.5, places=6)
            page.left_metric.setCurrentIndex(9)  # EE = 60 * (3.815 * VO2 + 1.232 * VCO2)
            self.assertAlmostEqual(page.compare_canvas.left_series[0][1][0], 159.52, places=2)

    def test_p6_experiment_time_switches_database_and_keeps_selection_on_refresh(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            paths = []
            export_paths = []
            export_root = os.path.join(temp_dir, "trajectory")
            for stamp, oxygen in (("2026_09_07_18_21_36_963", 21.0), ("2026_09_08_10_00_00_000", 19.0)):
                data_dir = os.path.join(temp_dir, f"experiment_{stamp}", "data")
                os.makedirs(data_dir)
                path = os.path.join(data_dir, "data.db")
                with sqlite3.connect(path) as connection:
                    connection.execute("CREATE TABLE Epoch_data_cage_1 (time INTEGER, oxygen_num REAL)")
                    connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (1, ?)", (oxygen,))
                paths.append(path)
                export_stamp = datetime.strptime(stamp, "%Y_%m_%d_%H_%M_%S_%f").strftime("%Y%m%d_%H%M%S_%f")[:-3]
                export_path = os.path.join(export_root, f"experiment_{export_stamp}_pid1")
                csv_dir = os.path.join(export_path, "cage_1", "data")
                os.makedirs(csv_dir)
                with open(os.path.join(csv_dir, "trajectory.csv"), "w", encoding="utf-8") as stream:
                    stream.write("captureTimestamp,status,X,Y\n1786071798,ok,3,4\n")
                export_paths.append(export_path)

            context = lambda: {
                **self.context(),
                "experiment_db_path": paths[0],
                "experiment_storage_root": temp_dir,
            }
            patcher = patch("public.component.workflow_pages.TRAJECTORY_EXPORT_DIR", export_root)
            patcher.start()
            self.addCleanup(patcher.stop)
            page = DataAnalysisPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            self.assertEqual(page.experiment_time_combo.currentData(), export_paths[0])
            self.assertEqual(page._compare_left_series[0][1], 21.0)
            self.assertEqual(page.left_metric.itemText(5), "跑轮")
            self.assertEqual(page.right_metric.itemText(5), "跑轮")

            page.experiment_time_combo.setCurrentIndex(page.experiment_time_combo.findData(export_paths[1]))
            self.assertEqual(page._compare_left_series[0][1], 19.0)
            self.assertEqual(page._comparison_export_rows()[0][0][1][0][1], 19.0)
            self.assertEqual(page._selected_experiment_health(), {})
            page.refresh_experiment_times_button.click()
            self.assertEqual(page.experiment_time_combo.currentData(), export_paths[1])

            page.experiment_time_combo.setCurrentIndex(0)
            self.assertEqual(page._compare_left_series, [])
            page.refresh_experiment_times_button.click()
            self.assertEqual(page.experiment_time_combo.currentIndex(), 0)

    def test_p6_history_lists_exported_excel_not_unexported_database_folders(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            exported = os.path.join(temp_dir, "1_2026_09_16_09_00_51_302.xlsx")
            with open(exported, "wb") as stream:
                stream.write(b"test listing only")
            for name in ("1_2026_09_16_09_00_51_302", "1_2026_09_15_14_15_52_117"):
                data_dir = os.path.join(temp_dir, name, "data")
                os.makedirs(data_dir)
                with sqlite3.connect(os.path.join(data_dir, "data.db")):
                    pass
            page = DataAnalysisPage(
                context_provider=lambda: {**self.context(), "experiment_storage_root": temp_dir},
                health_source=self.health,
            )
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            values = [page.experiment_time_combo.itemData(index) for index in range(page.experiment_time_combo.count())]
            self.assertIn(os.path.abspath(exported), values)
            self.assertFalse(any(str(value or "").endswith("data.db") for value in values))
            self.assertIn("Excel", page.experiment_time_combo.itemText(page.experiment_time_combo.findData(os.path.abspath(exported))))

    def test_p6_selected_excel_metric_does_not_mix_database_values(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            name = "1_2026_09_16_09_00_51_302"
            workbook_path = os.path.join(temp_dir, f"{name}.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "通道1"
            sheet.append(["获取时间", "氧浓度(%)"])
            sheet.append([datetime(2026, 9, 16, 9, 1, 0), 20.5])
            workbook.save(workbook_path)
            workbook.close()
            data_dir = os.path.join(temp_dir, name, "data")
            os.makedirs(data_dir)
            with sqlite3.connect(os.path.join(data_dir, "data.db")) as connection:
                connection.execute("CREATE TABLE Epoch_data_cage_1 (time TEXT, oxygen_num REAL)")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES ('2026-09-16 09:01:00', 99.0)")
            page = DataAnalysisPage(
                context_provider=lambda: {**self.context(), "experiment_storage_root": temp_dir},
                health_source=self.health,
            )
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            page.experiment_time_combo.setCurrentIndex(page.experiment_time_combo.findData(os.path.abspath(workbook_path)))
            self.assertEqual([value for _stamp, value in page._metric_channel_series({1}, 3)[0][1]], [20.5])
            self.assertEqual(page._metric_channel_series({1}, 4), [])

    def test_p6_bio_status_reads_selected_cage_latest_values(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                columns = "id INTEGER, WM_weight_num REAL, DWM_weight_num REAL, EM_weight_num REAL, ENM_running_wheel_num REAL"
                connection.execute(f"CREATE TABLE Epoch_data_cage_1 ({columns})")
                connection.execute(f"CREATE TABLE Epoch_data_cage_2 ({columns})")
                connection.execute("INSERT INTO Epoch_data_cage_1 VALUES (1, 23.2, 4.19, 7.44, 35769)")
                connection.execute("INSERT INTO Epoch_data_cage_2 VALUES (1, 25.4, 5.01, 8.20, 12345)")
                connection.commit()

            context = lambda: {
                **self.context(),
                "experiment_db_path": db_path,
                "channel_individuals": {
                    "CH01": {"id": "M001", "sex": "雄", "group": "对照组"},
                    "CH02": {"id": "M002", "sex": "雌", "group": "高脂组"},
                },
            }
            health = lambda: {
                "components": [{"scope": "CH01", "status": "normal"}],
                "alerts": [],
            }
            page = DataAnalysisPage(context_provider=context, health_source=health)
            self.addCleanup(page.deleteLater)
            page.refresh()
            self.assertIn("M001", page.bio_label.text())
            self.assertEqual(page.bio_metric_labels["weight"].text(), "23.2 g")
            self.assertEqual(page.bio_metric_labels["food"].text(), "4.19 g")
            self.assertEqual(page.bio_metric_labels["water"].text(), "7.44 mL")
            self.assertEqual(page.bio_metric_labels["movement"].text(), "35769 计数")
            self.assertEqual(page.bio_status_label.text(), "在线")

            page.channel_buttons[1].click()
            self.assertIn("M002", page.bio_label.text())
            self.assertEqual(page.bio_metric_labels["weight"].text(), "25.4 g")
            self.assertEqual(page.bio_metric_labels["water"].text(), "8.2 mL")

    def test_p6_trajectory_comparison_selects_channels_and_modules(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                for channel in (1, 2):
                    connection.execute(
                        f"CREATE TABLE DetectionResults_cage_{channel} (x REAL, y REAL, z REAL, time REAL)"
                    )
                    connection.executemany(
                        f"INSERT INTO DetectionResults_cage_{channel} VALUES (?, ?, ?, ?)",
                        [(0.0, 0.0, 0.0, 0.0), (0.01 * channel, 0.02, 0.003, 10.0), (0.02 * channel, 0.04, 0.006, 20.0)],
                    )
                connection.commit()

            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataAnalysisPage(context_provider=context, health_source=self.health)
            page.trajectory_channel_buttons[1].click()  # add CH02
            page.trajectory_module_buttons["xy"].click()
            self.assertEqual(page._trajectory_modules, {"xy"})
            self.assertTrue(page.trajectory_module_buttons["xy"].isChecked())
            self.assertFalse(page.trajectory_module_buttons["3d"].isChecked())
            self.assertFalse(page.trajectory_compare_cards["xy"].isHidden())
            self.assertTrue(page.trajectory_compare_cards["3d"].isHidden())
            page.trajectory_module_buttons["height"].click()
            page.trajectory_module_buttons["sleep"].click()
            page.trajectory_module_buttons["cumulative"].click()
            page.trajectory_module_buttons["total"].click()
            page.trajectory_module_buttons["minute"].click()
            page.refresh()
            self.assertEqual(page._trajectory_channels, {1, 2})
            self.assertEqual(len(page.trajectory_compare_plots["3d"].series), 2)
            self.assertEqual(len(page.trajectory_compare_plots["xy"].series), 2)
            self.assertEqual(len(page.trajectory_compare_plots["height"].series), 2)
            self.assertEqual(len(page.trajectory_compare_plots["total"].series), 2)
            self.assertGreater(page.trajectory_compare_plots["cumulative"].series[0][1][-1][1], 0.0)
            page.stop_refresh()
            page.deleteLater()
            self.app.processEvents()

    def test_p6_xy_uses_time_matched_legacy_csv_and_preserves_missing_z(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_dir = os.path.join(temp_dir, "monitor", "experiment_2026_08_07_11_03_13_461", "data")
            os.makedirs(db_dir)
            db_path = os.path.join(db_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute("CREATE TABLE DetectionResults_cage_1 (x REAL, y REAL, z REAL, time REAL)")
                connection.execute("INSERT INTO DetectionResults_cage_1 VALUES (50, 50, 50, 1)")
            export_root = os.path.join(temp_dir, "trajectory")
            csv_dir = os.path.join(export_root, "experiment_20260807_110313_461_pid1", "cage_1", "data")
            os.makedirs(csv_dir)
            with open(os.path.join(csv_dir, "trajectory.csv"), "w", encoding="utf-8") as stream:
                stream.write("timestamp,status,X,Y\n")
                stream.write("1786071798,ok,0,0\n")
                stream.write("1786071799,no_mouse,,\n")
                stream.write("1786071800,ok,3,4\n")
                stream.write("1786071801,ok,6,8\n")
            with patch("public.component.workflow_pages.TRAJECTORY_EXPORT_DIR", export_root):
                from datetime import datetime
                from Module.mouse_trajectory_analysis.analysis_core import find_matching_trajectory_export
                self.assertIsNotNone(find_matching_trajectory_export(export_root, datetime(2026, 8, 7, 11, 3, 13, 461000)))
                page = DataAnalysisPage(
                    context_provider=lambda: {**self.context(), "experiment_db_path": db_path},
                    health_source=self.health,
                )
                self.addCleanup(page.stop_refresh)
                self.addCleanup(page.deleteLater)
                page.trajectory_module_buttons["xy"].click()
                self.assertIsNotNone(
                    page._legacy_trajectory_analysis(),
                    f"selected={page.experiment_time_combo.currentData()!r}; db={page._selected_experiment_db_path!r}",
                )
                records = page._trajectory_records_for_channels()[1]
                self.assertEqual(len(records), 3)
                self.assertTrue(all(item["source"] == "legacy_trajectory_csv" for item in records))
                points, _elapsed, cumulative = page._trajectory_distance_metrics(records)
                self.assertEqual(points[1][:2], (3.0, 4.0))
                self.assertAlmostEqual(cumulative[-1], 5.0)
                output_path = os.path.join(temp_dir, "trajectory_export.csv")
                with patch("public.component.workflow_pages.QFileDialog.getSaveFileName", return_value=(output_path, "")):
                    page._export_trajectory_csv()
                with open(output_path, "r", encoding="utf-8-sig") as stream:
                    lines = stream.readlines()
                self.assertIn(",6.000000,8.000000,,", lines[-1])

    def test_p6_keeps_monitor_curve_excel_separate_from_trajectory_csv(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            monitor_root = os.path.join(temp_dir, "monitor")
            os.makedirs(monitor_root)
            workbook_path = os.path.join(monitor_root, "1_2026_08_07_11_03_13_461.xlsx")
            workbook = Workbook()
            monitor_sheet = workbook.active
            monitor_sheet.title = "通道1"
            monitor_sheet.append(["获取时间", "氧浓度(%)"])
            monitor_sheet.append([datetime(2026, 8, 7, 11, 3, 14), 20.5])
            # This sheet is intentionally coordinate-shaped. It must not be
            # mistaken for the trajectory source when a matching CSV exists.
            trajectory_looking_sheet = workbook.create_sheet("轨迹通道1")
            trajectory_looking_sheet.append(["获取时间", "X", "Y"])
            trajectory_looking_sheet.append([datetime(2026, 8, 7, 11, 3, 14), 99.0, 100.0])
            workbook.save(workbook_path)
            workbook.close()

            export_root = os.path.join(temp_dir, "trajectory")
            csv_dir = os.path.join(export_root, "experiment_20260807_110313_461_pid1", "cage_1", "data")
            os.makedirs(csv_dir)
            with open(os.path.join(csv_dir, "trajectory.csv"), "w", encoding="utf-8") as stream:
                stream.write("captureTimestamp,status,X,Y\n")
                stream.write("1786090994,ok,3,4\n")

            with patch("public.component.workflow_pages.TRAJECTORY_EXPORT_DIR", export_root):
                page = DataAnalysisPage(
                    context_provider=lambda: {**self.context(), "experiment_storage_root": monitor_root},
                    health_source=self.health,
                )
                self.addCleanup(page.stop_refresh)
                self.addCleanup(page.deleteLater)
                page.experiment_time_combo.setCurrentIndex(
                    page.experiment_time_combo.findData(os.path.abspath(workbook_path))
                )

                page.trajectory_module_buttons["xy"].click()
                analysis = page._legacy_trajectory_analysis()
                self.assertIsNotNone(analysis)
                self.assertEqual(
                    os.fspath(analysis.coordinate_source_path),
                    os.path.abspath(os.path.join(export_root, "experiment_20260807_110313_461_pid1")),
                )
                self.assertTrue(all("trajectory.csv" in source for source in analysis.channels[1].source_sheets))
                records = page._trajectory_records_for_channels()[1]
                self.assertEqual(records[0]["point"][:2], (3.0, 4.0))

                # The comparison curve still comes from the selected monitor
                # workbook, independently of the trajectory source above.
                monitor_series = page._metric_channel_series({1}, 3)
                self.assertEqual([value for _stamp, value in monitor_series[0][1]], [20.5])

    def test_p6_has_independent_trajectory_experiment_selector(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            export_root = os.path.join(temp_dir, "trajectory")
            export_paths = []
            for stamp in ("20260807_110313_461", "20260808_110313_461"):
                export_path = os.path.join(export_root, f"experiment_{stamp}_pid1")
                csv_dir = os.path.join(export_path, "cage_1", "data")
                os.makedirs(csv_dir)
                with open(os.path.join(csv_dir, "trajectory.csv"), "w", encoding="utf-8") as stream:
                    stream.write("captureTimestamp,status,X,Y\n1786090994,ok,3,4\n")
                export_paths.append(export_path)

            with patch("public.component.workflow_pages.TRAJECTORY_EXPORT_DIR", export_root):
                page = DataAnalysisPage(
                    context_provider=lambda: {**self.context(), "experiment_storage_root": temp_dir},
                    health_source=self.health,
                )
                self.addCleanup(page.stop_refresh)
                self.addCleanup(page.deleteLater)
                selected_index = page.trajectory_experiment_combo.findData(export_paths[1])
                self.assertGreaterEqual(selected_index, 0)
                page.trajectory_experiment_combo.setCurrentIndex(selected_index)
                self.assertEqual(page._selected_trajectory_export_path, export_paths[1])

                # Refreshing the monitor-experiment selector must not replace
                # the independently selected trajectory experiment.
                page.refresh_experiment_times_button.click()
                self.assertEqual(page.trajectory_experiment_combo.currentData(), export_paths[1])

    def test_p6_lists_legacy_trajectory_run_without_monitor_database(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            export_root = os.path.join(temp_dir, "trajectory")
            csv_dir = os.path.join(export_root, "experiment_20260807_110313_461_pid1", "cage_1", "data")
            os.makedirs(csv_dir)
            with open(os.path.join(csv_dir, "trajectory.csv"), "w", encoding="utf-8") as stream:
                stream.write("captureTimestamp,status,X,Y\n1786071798,ok,3,4\n")
            with patch("public.component.workflow_pages.TRAJECTORY_EXPORT_DIR", export_root):
                page = DataAnalysisPage(context_provider=self.context, health_source=self.health)
                self.addCleanup(page.stop_refresh)
                self.addCleanup(page.deleteLater)
                trajectory_path = os.path.dirname(os.path.dirname(csv_dir))
                index = page.experiment_time_combo.findData(trajectory_path)
                self.assertGreaterEqual(index, 0)
                page.experiment_time_combo.setCurrentIndex(index)
                page.trajectory_module_buttons["xy"].click()
                self.assertEqual(len(page._trajectory_records_for_channels()[1]), 1)

    def test_p6_behavior_recognition_is_a_comparison_module(self):
        page = DataAnalysisPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)
        self.assertIn("recognition", page.trajectory_module_buttons)
        page.trajectory_module_buttons["recognition"].click()
        self.assertEqual(page._trajectory_modules, {"recognition"})
        self.assertIs(page.trajectory_compare_plots["recognition"], page.analysis_behavior_chart)
        self.assertTrue(page.trajectory_compare_cards["recognition"].isVisible() or not page.isVisible())
        self.assertFalse(page.trajectory_play_button.isVisible())

    def test_p6_hidden_comparison_cards_survive_orphan_cleanup_and_refresh(self):
        page = DataAnalysisPage(context_provider=self.context, health_source=self.health)
        self.addCleanup(page.stop_refresh)
        self.addCleanup(page.deleteLater)
        for card in page.trajectory_compare_cards.values():
            self.assertIs(card.parentWidget(), page.trajectory_compare_container)
            if card.parentWidget() is None:  # mirrors main-window orphan cleanup
                card.deleteLater()
        self.app.processEvents()
        gc.collect()
        self.assertFalse(sip.isdeleted(page.analysis_recognition_channel_label))
        page.trajectory_module_buttons["recognition"].click()
        page.trajectory_module_buttons["xy"].click()
        page.refresh()  # the next two-second timer tick must remain safe
        self.assertIn("当前选择", page.analysis_recognition_channel_label.text())

    def test_p6_hidden_compatibility_canvases_survive_repeated_navigation(self):
        from index.MainWindow_index import MainWindow_Index

        class WindowHarness:
            _show_workflow_page = MainWindow_Index._show_workflow_page

            def __init__(self):
                self.tab_widget = QTabWidget()
                self.workflow_shell = Mock()

        window = WindowHarness()
        self.addCleanup(window.tab_widget.deleteLater)
        page = DataAnalysisPage(context_provider=self.context, health_source=self.health, parent=window.tab_widget)
        self.addCleanup(page.stop_refresh)
        self.assertIs(page._legacy_behaviour_widget.parentWidget(), page)
        self.assertIsNotNone(page.heatmap_canvas.parentWidget())
        self.assertIsNotNone(page.history_canvas.parentWidget())
        other = QWidget(window.tab_widget)
        for _ in range(3):
            window._show_workflow_page(page, 6)
            window._show_workflow_page(other, 5)
            self.app.processEvents()
        window._show_workflow_page(page, 6)
        self.assertFalse(sip.isdeleted(page.heatmap_canvas))
        self.assertFalse(sip.isdeleted(page.history_canvas))
        page.refresh()

    def test_orphan_window_cleanup_does_not_delete_qt_popup_frames(self):
        from index.MainWindow_index import MainWindow_Index

        popup = QFrame()
        popup.show()
        self.addCleanup(popup.close)
        harness = type("WindowHarness", (), {"modules": [], "tab_widget": QTabWidget()})()
        self.addCleanup(harness.tab_widget.deleteLater)
        fake_app = Mock()
        fake_app.topLevelWidgets.return_value = [popup]
        with patch("index.MainWindow_index.QApplication.instance", return_value=fake_app):
            MainWindow_Index._hide_legacy_module_windows(harness)
        self.assertTrue(popup.isVisible())

    def test_p6_infrared_temperature_uses_real_cage_samples_and_exports(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            db_path = os.path.join(temp_dir, "data.db")
            with sqlite3.connect(db_path) as connection:
                connection.execute(
                    "CREATE TABLE MouseInfrared_data_cage_1 (time TEXT, tmp_hs_max REAL)"
                )
                connection.executemany(
                    "INSERT INTO MouseInfrared_data_cage_1 VALUES (?, ?)",
                    [
                        ("2026-09-16 10:01:00", 37.2),
                        ("2026-09-16 10:00:00", 36.8),
                        ("2026-09-16 10:02:00", None),
                    ],
                )
                connection.execute(
                    "CREATE TABLE MouseInfrared_data_cage_2 (time TEXT, recognize_time TEXT, tmp_hs_mean REAL)"
                )
                connection.execute(
                    "INSERT INTO MouseInfrared_data_cage_2 VALUES (?, ?, ?)",
                    (None, "2026_09_16_10_00_30_000", 35.5),
                )
                connection.commit()

            context = lambda: {**self.context(), "experiment_db_path": db_path}
            page = DataAnalysisPage(context_provider=context, health_source=self.health)
            self.addCleanup(page.stop_refresh)
            self.addCleanup(page.deleteLater)
            page.trajectory_channel_buttons[1].click()
            page.trajectory_module_buttons["infrared"].click()

            self.assertEqual(page._trajectory_modules, {"infrared"})
            self.assertEqual(page.trajectory_compare_plots["infrared"].mode, "infrared")
            self.assertEqual(
                [value for _stamp, value in page.trajectory_compare_plots["infrared"].series[0][1]],
                [36.8, 37.2],
            )
            self.assertEqual(page.trajectory_compare_plots["infrared"].series[1][1][0][1], 35.5)
            self.assertIn("红外温度", page.export_trajectory_csv_button.text())
            self.assertFalse(page.trajectory_play_button.isVisible())

            output_path = os.path.join(temp_dir, "infrared.csv")
            with patch(
                "public.component.workflow_pages.QFileDialog.getSaveFileName",
                return_value=(output_path, "CSV 文件 (*.csv)"),
            ):
                page._export_trajectory_csv()
            with open(output_path, "r", encoding="utf-8-sig", newline="") as stream:
                exported = stream.read()
            self.assertIn("CH01,2026-09-16 10:00:00,36.8", exported)
            self.assertIn("CH02,2026_09_16_10_00_30_000,35.5", exported)
            self.assertNotIn("10:02:00", exported)

            plot = page.trajectory_compare_plots["infrared"]
            plot.resize(640, 340)
            image = QImage(640, 340, QImage.Format.Format_ARGB32)
            painter = QPainter(image)
            plot.render(painter)
            self.assertTrue(painter.isActive())
            painter.end()

    def test_p6_trajectory_comparison_canvas_renders_without_active_painter(self):
        plot = TrajectoryComparisonPlot("3d")
        plot.resize(640, 320)
        plot.set_series([("CH01", [(0.0, 0.0, 0.0), (0.01, 0.02, 0.003)])])
        image = QImage(640, 320, QImage.Format.Format_ARGB32)
        painter = QPainter(image)
        plot.render(painter)
        self.assertTrue(painter.isActive())
        painter.end()
        plot.deleteLater()
        self.app.processEvents()


if __name__ == "__main__":
    unittest.main()
